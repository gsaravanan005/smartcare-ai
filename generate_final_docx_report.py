"""
SmartCare AI - Final Year Project Report Generator: Front Matter & Meta
Configured with user's institutional details:
- Rajalakshmi Engineering College (Autonomous), Chennai-602 105
- Ms. SUBASHREE R M.Tech., Supervisor & Assistant Professor
- Dr. P. VALARMATHIE M.E., Ph.D., Professor & Head, IT Department
- SANGARI A C (231001178) & SARAVANAN G (231001504)
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from report_builder_core import (
    set_cell_background, set_cell_margins, set_table_borders,
    add_callout_box, add_styled_paragraph, add_heading_1,
    add_heading_2, add_heading_3, add_bullet_item
)

PROJECT_TITLE = "SMARTCARE AI: COMORBIDITY-AWARE EXPLAINABLE MULTI-TASK LEARNING FOR INTEGRATED DIABETES, CARDIOVASCULAR, AND CHRONIC KIDNEY DISEASE RISK PREDICTION"
STUDENT_1 = "SANGARI A C (231001178)"
STUDENT_2 = "SARAVANAN G (231001504)"
SUPERVISOR_NAME = "Ms. SUBASHREE R M.Tech."
HOD_NAME = "Dr. P. VALARMATHIE M.E., Ph.D."
COLLEGE_NAME = "RAJALAKSHMI ENGINEERING COLLEGE"
COLLEGE_SUB = "(An Autonomous Institution Affiliated to Anna University Chennai)"
COLLEGE_LOC = "(AUTONOMOUS), CHENNAI-602 105"
DEPT_NAME = "DEPARTMENT OF INFORMATION TECHNOLOGY"
DEGREE_TEXT = "BACHELOR OF TECHNOLOGY\nin\nINFORMATION TECHNOLOGY"
MONTH_YEAR = "APRIL 2026"

def create_full_report():
    doc = Document()
    
    # Page Setup - Standard 1 inch / 1.25 inch academic margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)  # A4
        section.page_height = Inches(11.69)
        
        # Header & Footer setup
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("SmartCare AI: Comorbidity-Aware Multi-Task Clinical Decision Support System")
        r_ft.font.name = "Times New Roman"
        r_ft.font.size = Pt(8.5)
        r_ft.font.italic = True
        r_ft.font.color.rgb = RGBColor(100, 116, 139)

    # -------------------------------------------------------------
    # PAGE 1: TITLE / COVER PAGE
    # -------------------------------------------------------------
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(36)
    p_t.paragraph_format.space_after = Pt(12)
    p_t.paragraph_format.line_spacing = 1.3
    r = p_t.add_run(f"{PROJECT_TITLE}\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(8)
    p_sub.paragraph_format.space_after = Pt(20)
    r = p_sub.add_run("IT19811 PROJECT PHASE-II REPORT\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(3, 105, 161)
    
    p_by = doc.add_paragraph()
    p_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_by.paragraph_format.space_before = Pt(12)
    p_by.paragraph_format.space_after = Pt(8)
    r = p_by.add_run(f"Submitted by\n\n{STUDENT_1}\n{STUDENT_2}\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.space_before = Pt(16)
    p_deg.paragraph_format.space_after = Pt(20)
    p_deg.paragraph_format.line_spacing = 1.3
    r = p_deg.add_run(f"in partial fulfilment for the award of the degree of\n{DEGREE_TEXT}\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    add_callout_box(doc, "[INSERT RAJALAKSHMI ENGINEERING COLLEGE EMBLEM / LOGO HERE]", box_type="placeholder")
    
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(16)
    p_inst.paragraph_format.space_after = Pt(4)
    p_inst.paragraph_format.line_spacing = 1.2
    r = p_inst.add_run(f"{DEPT_NAME}\n{COLLEGE_NAME}\n{COLLEGE_LOC}\n\n{MONTH_YEAR}")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 2: BONAFIDE CERTIFICATE
    # -------------------------------------------------------------
    p_c1 = doc.add_paragraph()
    p_c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c1.paragraph_format.space_before = Pt(10)
    p_c1.paragraph_format.space_after = Pt(4)
    r = p_c1.add_run(f"{COLLEGE_NAME}\n{COLLEGE_SUB}")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    p_cert = doc.add_paragraph()
    p_cert.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert.paragraph_format.space_before = Pt(16)
    p_cert.paragraph_format.space_after = Pt(16)
    r = p_cert.add_run("BONAFIDE CERTIFICATE")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    
    add_styled_paragraph(
        doc,
        f"Certified that this Phase-II Thesis titled “{PROJECT_TITLE.title()}” is the Bonafide work of {STUDENT_1}, {STUDENT_2} who carried out the work under my supervision. Certified further that to the best of my knowledge the work reported herein does not form part of any other thesis or dissertation on the basis of which a degree or award was conferred on an earlier occasion on this or any other candidate.",
        space_after=14, line_spacing=1.3
    )
    
    p_sdg = doc.add_paragraph()
    p_sdg.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_sdg.paragraph_format.space_before = Pt(8)
    p_sdg.paragraph_format.space_after = Pt(36)
    r = p_sdg.add_run("“This project addresses the following Sustainable Development Goals:\nSDG 3, SDG 8, SDG 9, SDG 10 & SDG 12”.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.italic = True
    
    # Signature Table
    tbl_sig = doc.add_table(rows=2, cols=2)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sig.autofit = False
    tbl_sig.columns[0].width = Inches(3.2)
    tbl_sig.columns[1].width = Inches(3.2)
    
    c0 = tbl_sig.cell(0, 0)
    p = c0.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(f"{HOD_NAME}\nProfessor & Head\nDepartment of Information Technology\nRajalakshmi Engineering College\nChennai- 602105")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.bold = True
    
    c1 = tbl_sig.cell(0, 1)
    p = c1.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(f"{SUPERVISOR_NAME}\nSupervisor & Assistant Professor\nDepartment of Information Technology\nRajalakshmi Engineering College\nChennai- 602105")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.bold = True
    
    p_viva = doc.add_paragraph()
    p_viva.paragraph_format.space_before = Pt(36)
    p_viva.paragraph_format.space_after = Pt(36)
    r = p_viva.add_run("Submitted to Project Viva-Voce Examination held on ............................................")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    
    tbl_ex = doc.add_table(rows=1, cols=2)
    tbl_ex.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ex.columns[0].width = Inches(3.2)
    tbl_ex.columns[1].width = Inches(3.2)
    
    p = tbl_ex.cell(0, 0).paragraphs[0]
    r = p.add_run("Internal Examiner")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    
    p = tbl_ex.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("External Examiner")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 3: ACKNOWLEDGEMENT
    # -------------------------------------------------------------
    p_ack = doc.add_paragraph()
    p_ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ack.paragraph_format.space_before = Pt(10)
    p_ack.paragraph_format.space_after = Pt(16)
    r = p_ack.add_run("ACKNOWLEDGEMENT")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    
    add_styled_paragraph(
        doc,
        "First, we thank the almighty God for the successful completion of the project. Our heartfelt sincere thanks to our beloved chairman Mr. S. Meganathan B.E., F.I.E., for his sincere endeavor in educating us in his premier institution. We would like to express our deep gratitude to our beloved Chairperson Dr. Thangam Meganathan Ph.D., for her enthusiastic motivation which inspired us a lot in completing this project and Vice-Chairman Mr. Abhay Shankar Meganathan B.E., M.S., for providing us with the requisite infrastructure.",
        space_after=12, line_spacing=1.3
    )
    
    add_styled_paragraph(
        doc,
        "We also express our sincere gratitude to our college principal, Dr. S. N. Murugesan M.E., Ph.D., for his kind support and facilities to complete our work on time. We extend heartfelt gratitude to Dr. P. Valarmathie M.E., Ph.D., Professor and Head of the Department of Information Technology for her guidance and encouragement throughout the work. We are very glad to thank our project coordinator Dr. M. Babu M.E., Ph.D., Professor for his encouragement and support towards the successful completion of this project.",
        space_after=12, line_spacing=1.3
    )
    
    add_styled_paragraph(
        doc,
        "Further We express our deepest gratitude to our supervisor Ms. Subashree R M.Tech., Assistant Professor for her valuable guidance throughout the work. We extend our thanks to our parents, friends, all faculty members, and supporting staff for their direct and indirect involvement in the successful completion of the project for their encouragement and support.",
        space_after=24, line_spacing=1.3
    )
    
    p_names = doc.add_paragraph()
    p_names.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_names.paragraph_format.line_spacing = 1.2
    r = p_names.add_run("Sangari A C\nSaravanan G")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 4: ABSTRACT
    # -------------------------------------------------------------
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs.paragraph_format.space_before = Pt(10)
    p_abs.paragraph_format.space_after = Pt(16)
    r = p_abs.add_run("ABSTRACT")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    
    add_styled_paragraph(
        doc,
        "The “SmartCare AI: Comorbidity-Aware Explainable Multi-Task Learning for Integrated Diabetes, Cardiovascular, and Chronic Kidney Disease Risk Prediction” research tackles the growing global burden of chronic non-communicable diseases through unified multi-task machine learning. In clinical reality, Type 2 Diabetes (T2D), Cardiovascular Disease (CVD), and Chronic Kidney Disease (CKD) exhibit intricate pathophysiological interactions, shared metabolic risk factors, and compounding vascular complications. Conventional single-disease models operate in silos as opaque 'black boxes' and fail to capture these multidirectional dependencies. SmartCare AI introduces a unified multi-task neural network with a shared 32-dimensional latent feature space that simultaneously predicts calibrated risk probabilities for T2D, CVD, and CKD from a harmonized 30-dimensional clinical feature vector.",
        space_after=12, line_spacing=1.3
    )
    
    add_styled_paragraph(
        doc,
        "The system harmonizes three major clinical cohorts—CDC BRFSS 2015 (N=70,692), Kaggle CVD (N=68,205), and UCI CKD (N=400)—using zero-leakage median imputation and RobustScaler transformations. The multi-task network is trained with task-masked loss and task-balanced batch sampling (20% CKD allocation), followed by post-hoc Platt sigmoid scaling. On strictly held-out test partitions, SmartCare AI achieves high diagnostic precision and sensitivity: T2D (ROC-AUC: 0.8276, Recall: 88.72%, F1: 0.7751), CVD (ROC-AUC: 0.7974, Recall: 79.13%, F1: 0.7306), and CKD (ROC-AUC: 0.9953, Recall: 94.59%, F1: 0.9589) with excellent calibration (ECE < 0.011). Game-theoretic SHAP explainability provides local waterfall attributions and global biomarker importance, while a deterministic clinical rules engine translates risks into ADA/ACC/KDIGO-guideline-adherent care plans within a FastAPI and React microservices architecture.",
        space_after=12, line_spacing=1.3
    )
    
    add_styled_paragraph(
        doc,
        "The research contributes directly to Sustainable Development Goals 3 (Good Health and Well-Being), 8 (Decent Work and Economic Growth), 9 (Industry, Innovation, and Infrastructure), 10 (Reduced Inequalities), and 12 (Responsible Consumption and Production) by enabling early, equitable, and explainable chronic disease risk stratification. By replacing disjointed single-disease calculators with a unified comorbidity framework, SmartCare AI empowers clinicians and patients with actionable, data-driven preventative intelligence for improved clinical outcomes and healthcare sustainability.",
        space_after=16, line_spacing=1.3
    )
    
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGES 5-6: TABLE OF CONTENTS
    # -------------------------------------------------------------
    p_toc = doc.add_paragraph()
    p_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc.paragraph_format.space_before = Pt(10)
    p_toc.paragraph_format.space_after = Pt(16)
    r = p_toc.add_run("TABLE OF CONTENTS")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    
    toc_data = [
        ("CHAPTER NO", "TITLE", "PAGE NO"),
        ("", "ABSTRACT", "iv"),
        ("", "LIST OF FIGURES", "vii"),
        ("", "LIST OF TABLES", "viii"),
        ("", "LIST OF ABBREVIATIONS", "ix"),
        ("", "DEPARTMENT VISION", "x"),
        ("", "DEPARTMENT MISSION", "x"),
        ("", "PROGRAMME EDUCATIONAL OBJECTIVES (PEOs)", "x"),
        ("", "PROGRAM OUTCOMES (POs)", "xi"),
        ("", "PROGRAM SPECIFIC OUTCOMES (PSOs)", "xii"),
        ("", "COURSE OBJECTIVE & COURSE OUTCOMES", "xiii"),
        ("1", "INTRODUCTION", "1"),
        ("", "1.1 MOTIVATION", "1"),
        ("", "1.2 EXISTING SYSTEM", "2"),
        ("", "1.3 PROBLEM STATEMENT", "5"),
        ("", "1.4 OBJECTIVES OF THE PROJECT", "6"),
        ("", "1.5 PROPOSED SYSTEM", "8"),
        ("", "1.6 BENEFITS OF THE PROJECT", "9"),
        ("2", "LITERATURE SURVEY", "11"),
        ("", "2.1 INTRODUCTION", "11"),
        ("", "2.2 RELATED WORK", "12"),
        ("", "2.3 INFERENCE FROM RELATED WORK", "16"),
        ("3", "SYSTEM DESIGN", "17"),
        ("", "3.1 INTRODUCTION", "17"),
        ("", "3.2 SYSTEM ARCHITECTURE", "18"),
        ("", "3.3 SYSTEM REQUIREMENTS", "21"),
        ("", "3.4 DATA FLOW DIAGRAM / USE CASE DIAGRAM", "25"),
        ("4", "PROJECT DESCRIPTION", "29"),
        ("", "4.1 METHODOLOGIES", "29"),
        ("", "4.2 MODULE DESCRIPTION", "31"),
        ("5", "RESULT AND DISCUSSION", "41"),
        ("6", "CONCLUSION AND FUTURE WORK", "49"),
        ("", "REFERENCES", "51"),
        ("", "APPENDICES", "53"),
        ("I", "PROJECT PLAGIARISM REPORT", "53"),
        ("II", "PAPER PLAGIARISM REPORT", "54"),
        ("IV", "PROOF OF THE PUBLICATION", "61"),
        ("V", "CO-PO-PSO MAPPING", "64"),
        ("VI", "CO-SDG RELEVANCE RECORD", "65"),
    ]
    
    tbl_toc = doc.add_table(rows=len(toc_data), cols=3)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_toc.autofit = False
    tbl_toc.columns[0].width = Inches(1.5)
    tbl_toc.columns[1].width = Inches(4.2)
    tbl_toc.columns[2].width = Inches(1.0)
    
    for idx, (c1_txt, c2_txt, c3_txt) in enumerate(toc_data):
        row = tbl_toc.rows[idx]
        for col_i, txt in enumerate([c1_txt, c2_txt, c3_txt]):
            cell = row.cells[col_i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if col_i == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif col_i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            r = p.add_run(txt)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5 if idx > 0 else 11)
            if idx == 0 or c1_txt in ["1", "2", "3", "4", "5", "6", "REFERENCES", "APPENDICES"] or txt.startswith("CHAPTER"):
                r.font.bold = True
                
        if idx == 0:
            for cell in row.cells:
                set_cell_background(cell, "F1F5F9")
                
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 7: LIST OF FIGURES
    # -------------------------------------------------------------
    p_lof = doc.add_paragraph()
    p_lof.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lof.paragraph_format.space_before = Pt(10)
    p_lof.paragraph_format.space_after = Pt(16)
    r = p_lof.add_run("LIST OF FIGURES")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    
    figures_data = [
        ("Figure Number", "Figure Caption", "Page Number"),
        ("3.1", "System Architecture Diagram", "19"),
        ("3.2", "Data Flow Diagram", "25"),
        ("3.3", "Use Case Diagram", "28"),
        ("4.1", "Methodology", "30"),
        ("5.1", "Multi-Task Neural Network Training & Loss Convergence", "42"),
        ("5.2", "Performance Evaluation", "43"),
        ("5.3", "Recommendation Module (Comorbidity Risk Dashboard)", "43"),
        ("5.4", "SHAP Biomarker Attribution & Risk Decomposition", "44"),
        ("5.5", "Clinical Guideline Intervention Plan", "44"),
        ("5.6", "Longitudinal Health Trends Analysis", "45"),
        ("5.7", "Emergency Alert Dispatcher & Clinical Telemetry", "46"),
        ("5.8", "Multi-Lingual Localization Interface", "47"),
        ("5.9", "AI Assistance Chatbot (Clinical Copilot)", "48"),
    ]
    
    tbl_lof = doc.add_table(rows=len(figures_data), cols=3)
    tbl_lof.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_lof.autofit = False
    tbl_lof.columns[0].width = Inches(1.5)
    tbl_lof.columns[1].width = Inches(4.2)
    tbl_lof.columns[2].width = Inches(1.0)
    
    for idx, (c1_txt, c2_txt, c3_txt) in enumerate(figures_data):
        row = tbl_lof.rows[idx]
        for col_i, txt in enumerate([c1_txt, c2_txt, c3_txt]):
            cell = row.cells[col_i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if col_i == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif col_i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            r = p.add_run(txt)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5 if idx > 0 else 11)
            if idx == 0:
                r.font.bold = True
                set_cell_background(cell, "F1F5F9")
                
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 8: LIST OF TABLES
    # -------------------------------------------------------------
    p_lot = doc.add_paragraph()
    p_lot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lot.paragraph_format.space_before = Pt(10)
    p_lot.paragraph_format.space_after = Pt(16)
    r = p_lot.add_run("LIST OF TABLES")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    
    tables_data = [
        ("Table Number", "Table Caption", "Page Number"),
        ("1.1", "Existing Vs Proposed System", "4"),
        ("3.1", "Software Tools", "22"),
        ("3.2", "Functional Requirements", "23"),
        ("3.3", "Hardware Requirements", "24"),
        ("4.1", "30-Dimensional Harmonized Clinical Feature Space", "30"),
        ("5.1", "Multi-Task Neural Network Test Set Performance Metrics", "41"),
        ("5.2", "Empirical Baseline Comparison", "42"),
        ("5.3", "Ablation Study on Shared Representation & Task Loss Weighting", "43"),
    ]
    
    tbl_lot = doc.add_table(rows=len(tables_data), cols=3)
    tbl_lot.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_lot.autofit = False
    tbl_lot.columns[0].width = Inches(1.5)
    tbl_lot.columns[1].width = Inches(4.2)
    tbl_lot.columns[2].width = Inches(1.0)
    
    for idx, (c1_txt, c2_txt, c3_txt) in enumerate(tables_data):
        row = tbl_lot.rows[idx]
        for col_i, txt in enumerate([c1_txt, c2_txt, c3_txt]):
            cell = row.cells[col_i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if col_i == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif col_i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            r = p.add_run(txt)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5 if idx > 0 else 11)
            if idx == 0:
                r.font.bold = True
                set_cell_background(cell, "F1F5F9")
                
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGE 9: LIST OF ABBREVIATIONS
    # -------------------------------------------------------------
    p_loa = doc.add_paragraph()
    p_loa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_loa.paragraph_format.space_before = Pt(10)
    p_loa.paragraph_format.space_after = Pt(16)
    r = p_loa.add_run("LIST OF ABBREVIATIONS")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    
    abbr_data = [
        ("S.No", "Abbreviation", "Word Expansion"),
        ("1", "AI", "Artificial Intelligence"),
        ("2", "ML", "Machine Learning"),
        ("3", "MTL", "Multi-Task Learning"),
        ("4", "T2D", "Type 2 Diabetes Mellitus"),
        ("5", "CVD", "Cardiovascular Disease"),
        ("6", "CKD", "Chronic Kidney Disease"),
        ("7", "SHAP", "SHapley Additive exPlanations"),
        ("8", "ECE", "Expected Calibration Error"),
        ("9", "ROC-AUC", "Receiver Operating Characteristic - Area Under Curve"),
        ("10", "PR-AUC", "Precision-Recall - Area Under Curve"),
        ("11", "BMI", "Body Mass Index"),
        ("12", "SBP", "Systolic Blood Pressure"),
        ("13", "DBP", "Diastolic Blood Pressure"),
        ("14", "eGFR", "Estimated Glomerular Filtration Rate"),
        ("15", "FBS", "Fasting Blood Sugar"),
        ("16", "HbA1c", "Glycated Hemoglobin"),
        ("17", "RBAC", "Role-Based Access Control"),
        ("18", "REST", "Representational State Transfer"),
        ("19", "API", "Application Programming Interface"),
        ("20", "JWT", "JSON Web Token"),
        ("21", "BRFSS", "Behavioral Risk Factor Surveillance System (CDC)"),
        ("22", "UCI", "University of California, Irvine"),
        ("23", "WHO", "World Health Organization"),
        ("24", "ADA", "American Diabetes Association"),
        ("25", "ACC/AHA", "American College of Cardiology / American Heart Association"),
        ("26", "KDIGO", "Kidney Disease: Improving Global Outcomes"),
        ("27", "SDG", "Sustainable Development Goals"),
        ("28", "LLM", "Large Language Model"),
        ("29", "XAI", "Explainable Artificial Intelligence"),
        ("30", "EHR", "Electronic Health Record"),
    ]
    
    tbl_loa = doc.add_table(rows=len(abbr_data), cols=3)
    tbl_loa.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_loa.autofit = False
    tbl_loa.columns[0].width = Inches(1.0)
    tbl_loa.columns[1].width = Inches(2.2)
    tbl_loa.columns[2].width = Inches(3.5)
    
    for idx, (c1_txt, c2_txt, c3_txt) in enumerate(abbr_data):
        row = tbl_loa.rows[idx]
        for col_i, txt in enumerate([c1_txt, c2_txt, c3_txt]):
            cell = row.cells[col_i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if col_i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            r = p.add_run(txt)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5 if idx > 0 else 11)
            if idx == 0:
                r.font.bold = True
                set_cell_background(cell, "F1F5F9")
                
    doc.add_page_break()

    # -------------------------------------------------------------
    # PAGES 10-13: REC DEPARTMENT VISION, MISSION, PEOs, POs, PSOs, COs
    # -------------------------------------------------------------
    p_vis = doc.add_paragraph()
    p_vis.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_vis.paragraph_format.space_before = Pt(6)
    p_vis.paragraph_format.space_after = Pt(4)
    r = p_vis.add_run("DEPARTMENT VISION")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    
    add_styled_paragraph(doc, "To be a Department of Excellence in Information Technology Education, Research and Development.", align=WD_ALIGN_PARAGRAPH.CENTER)
    
    p_mis = doc.add_paragraph()
    p_mis.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mis.paragraph_format.space_before = Pt(8)
    p_mis.paragraph_format.space_after = Pt(4)
    r = p_mis.add_run("DEPARTMENT MISSION")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    
    add_bullet_item(doc, "M1", "To train the students to become highly knowledgeable in the field of Information Technology.")
    add_bullet_item(doc, "M2", "To promote continuous learning and research in core and emerging areas.")
    add_bullet_item(doc, "M3", "To develop globally competent students with strong foundations, who will be able to adapt to changing technologies.")

    p_peo = doc.add_paragraph()
    p_peo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_peo.paragraph_format.space_before = Pt(12)
    p_peo.paragraph_format.space_after = Pt(4)
    r = p_peo.add_run("PROGRAMME EDUCATIONAL OBJECTIVES")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    
    add_bullet_item(doc, "PEO I", "To provide essential background in Science, basic Electronics and applied Mathematics.")
    add_bullet_item(doc, "PEO II", "To prepare the students with fundamental knowledge in programming languages and to develop applications.")
    add_bullet_item(doc, "PEO III", "To engage the students in life-long learning, and make them to remain current in their profession and obtain additional qualifications to enhance their career positions in IT industries.")
    add_bullet_item(doc, "PEO IV", "To enable the students to implement computing solutions for real world problems and carry out basic and applied research leading to new innovations in Information Technology (IT) and related interdisciplinary areas.")
    add_bullet_item(doc, "PEO V", "To familiarize the students with the ethical issues in engineering profession, issues related to the world wide economy, nurturing of current job related skills and emerging technologies.")

    doc.add_page_break()

    # POs Page
    p_po = doc.add_paragraph()
    p_po.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_po.paragraph_format.space_before = Pt(6)
    p_po.paragraph_format.space_after = Pt(8)
    r = p_po.add_run("PROGRAM OUTCOMES (POs)")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    
    pos = [
        ("1. Engineering knowledge", "Apply the knowledge of mathematics, science, engineering fundamentals, and an engineering specialization to the solution of complex engineering problems."),
        ("2. Problem analysis", "Identify, formulate, review research literature, and analyze complex engineering problems reaching substantiated conclusions using first principles of mathematics, natural sciences, and engineering sciences."),
        ("3. Design/development of solutions", "Design solutions for complex engineering problems and design system components or processes that meet the specified needs with appropriate consideration for the public health and safety, and the cultural, societal, and environmental considerations."),
        ("4. Conduct investigations of complex problems", "Use research-based knowledge and research methods including design of experiments, analysis and interpretation of data, and synthesis of the information to provide valid conclusions."),
        ("5. Modern tool usage", "Create, select, and apply appropriate techniques, resources, and modern engineering and IT tools including prediction and modeling to complex engineering activities with an understanding of the limitations."),
        ("6. The engineer and society", "Apply reasoning informed by the contextual knowledge to assess societal, health, safety, legal and cultural issues and the consequent responsibilities relevant to the professional engineering practice."),
        ("7. Environment and sustainability", "Understand the impact of the professional engineering solutions in societal and environmental contexts, and demonstrate the knowledge of, and need for sustainable development."),
        ("8. Ethics", "Apply ethical principles and commit to professional ethics and responsibilities and norms of the engineering practice."),
        ("9. Individual and team work", "Function effectively as an individual, and as a member or leader in diverse teams, and in multidisciplinary settings."),
        ("10. Communication", "Communicate effectively on complex engineering activities with the engineering community and with society at large, such as, being able to comprehend and write effective reports and design documentation, make effective presentations, and give and receive clear instructions."),
        ("11. Project management and finance", "Demonstrate knowledge and understanding of the engineering and management principles and apply these to one’s own work, as a member and leader in a team, to manage projects and in multidisciplinary environments."),
        ("12. Life-long learning", "Recognize the need for and have the preparation and ability to engage in independent and life-long learning in the broadest context of technological change.")
    ]
    for title, desc in pos:
        add_bullet_item(doc, title, desc)

    doc.add_page_break()

    # PSOs and COs Page
    p_pso = doc.add_paragraph()
    p_pso.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pso.paragraph_format.space_before = Pt(6)
    p_pso.paragraph_format.space_after = Pt(6)
    r = p_pso.add_run("PROGRAM SPECIFIC OUTCOMES (PSOs)")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    
    psos = [
        ("1", "To comprehend and analyze user requirements to design IT based solutions."),
        ("2", "To identify and assess current technologies and review their applicability to address individual and organizational needs."),
        ("3", "To engage in the computing profession by working effectively and utilizing professional skills to make a positive contribution to society."),
        ("4", "To take on positions as promoters in business and embark on a research career in the field.")
    ]
    for title, desc in psos:
        add_bullet_item(doc, title, desc)

    p_co = doc.add_paragraph()
    p_co.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_co.paragraph_format.space_before = Pt(12)
    p_co.paragraph_format.space_after = Pt(6)
    r = p_co.add_run("COURSE OBJECTIVE")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    
    add_styled_paragraph(doc, "• To enable the students to do the implementation of the industry-relevant and real-time projects on various core domains of information technology.", space_after=8)

    p_co_head = doc.add_paragraph()
    p_co_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_co_head.paragraph_format.space_before = Pt(12)
    p_co_head.paragraph_format.space_after = Pt(6)
    r = p_co_head.add_run("COURSE OUTCOMES")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    
    cos = [
        ("•", "Analyze complex Engineering problems related Information Technology to reach substantiated conclusions by applying knowledge of Mathematics, Engineering fundamentals and Engineering specialization."),
        ("•", "Create research based solutions for complex computer Engineering or multidisciplinary problems, and design system components or processes by applying appropriate techniques, resources, and modern IT tools."),
        ("•", "Apply contextual computer science engineering solutions in the sustainable development towards environmental, societal, health, safety, legal, cultural issues and needs."),
        ("•", "Apply ethical principles and commit to professional ethics and responsibilities and norms of the engineering practice."),
        ("•", "Perform effectively as an individual, and as a member or leader in diverse teams, Communicate effectively and write effective reports and design documentation, ability to engage themselves in life-long learning.")
    ]
    for bullet, desc in cos:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(f"{bullet} {desc}")
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)

    doc.add_page_break()

    return doc
