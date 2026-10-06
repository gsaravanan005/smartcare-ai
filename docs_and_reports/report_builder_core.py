"""
SmartCare AI - Core Formatting and Figure Integration Utilities
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Applies borders to the entire table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def add_callout_box(doc, text_content, box_type="placeholder"):
    """Creates a stylized placeholder or callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.0)
    
    if box_type == "placeholder":
        set_cell_background(cell, "F8FAFC")
        border_xml = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="dashed" w:sz="12" w:space="0" w:color="0284C7"/><w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="0284C7"/><w:left w:val="dashed" w:sz="12" w:space="0" w:color="0284C7"/><w:right w:val="dashed" w:sz="12" w:space="0" w:color="0284C7"/></w:tcBorders>')
    else:
        set_cell_background(cell, "F1F5F9")
        border_xml = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:left w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/></w:tcBorders>')
        
    cell._tc.get_or_add_tcPr().append(border_xml)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text_content)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor(2, 132, 199) if box_type == "placeholder" else RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(4)

def add_figure_with_image(doc, image_path, caption_text, width_inches=5.8, space_before=8, space_after=12):
    """Inserts a high-resolution figure image with centered academic caption."""
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(space_before)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(space_after)
        p_cap.paragraph_format.keep_with_next = False
        r = p_cap.add_run(caption_text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.italic = True
        return p_cap
    else:
        add_callout_box(doc, f"[INSERT: {caption_text}]", box_type="placeholder")
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(space_after)
        r = p_cap.add_run(caption_text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.italic = True
        return p_cap

def add_styled_paragraph(doc, text, style='Body', bold_prefix=None, space_after=8, line_spacing=1.35, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Adds a formatted paragraph with optional bold prefix and typography controls."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
        
    if text:
        r_text = p.add_run(text)
        r_text.font.name = 'Times New Roman'
        r_text.font.size = Pt(12)
        r_text.font.color.rgb = RGBColor(30, 41, 59)
        
    return p

def add_heading_1(doc, text):
    """Adds Chapter/Major Section heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_heading_2(doc, text):
    """Adds Subsection heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(3, 105, 161)
    return p

def add_heading_3(doc, text):
    """Adds Sub-subsection heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_bullet_item(doc, title, description):
    """Adds a structured bullet point."""
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    
    r_bold = p.add_run(f"{title}: ")
    r_bold.font.name = 'Times New Roman'
    r_bold.font.size = Pt(12)
    r_bold.font.bold = True
    r_bold.font.color.rgb = RGBColor(15, 23, 42)
    
    r_desc = p.add_run(description)
    r_desc.font.name = 'Times New Roman'
    r_desc.font.size = Pt(12)
    r_desc.font.color.rgb = RGBColor(30, 41, 59)
