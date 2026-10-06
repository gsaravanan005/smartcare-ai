"""
SmartCare AI - Final Year Project Report Generator
Generates a complete, comprehensive, academic-grade DOCX report matching the 78-page sample reference structure.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_callout_box(doc, text_content, box_type="placeholder"):
    """Creates a stylized placeholder or callout box for figures/screenshots."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.2)
    
    if box_type == "placeholder":
        set_cell_background(cell, "F8FAFC")
        border_xml = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="dashed" w:sz="12" w:space="0" w:color="0284C7"/><w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="0284C7"/><w:left w:val="dashed" w:sz="12" w:space="0" w:color="0284C7"/><w:right w:val="dashed" w:sz="12" w:space="0" w:color="0284C7"/></w:tcBorders>')
    else:
        set_cell_background(cell, "F1F5F9")
        border_xml = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:left w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/></w:tcBorders>')
        
    cell._tc.get_or_add_tcPr().append(border_xml)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text_content)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor(2, 132, 199) if box_type == "placeholder" else RGBColor(30, 41, 59)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Helper functions defined successfully.")
